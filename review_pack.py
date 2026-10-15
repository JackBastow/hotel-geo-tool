"""
The review pack: the hotel uploads the review export it already has (Booking.com extranet, Tripadvisor, Google Business Profile, Expedia Partner Central, or any
spreadsheet with a review and, where it exists, a reply), and the audit analyses the WHOLE file instead of the handful of reviews public search happens to show.

Why this route: the platforms prohibit automated collection and this project has no paid APIs, but every platform lets the hotel export its own reviews. The hotel
owns that data and chooses to share it.

What it works out (plain counting and rules; nothing here asks an AI anything):
  * reply rate overall, by platform, by rating band and over the last 90 days of the file; critical reviews still unanswered;
  * response time where both dates exist;
  * how the replies were written: specific to the guest, thanks or apologises, says what will be done, signed by a role or team;
  * how many replies are identical or near-identical templates;
  * what guests praise and what they criticise, by topic, set against what the hotel's own site states;
  * rating and volume by month, and how recent the reviews are.

Privacy: reviews carry personal data. analyse() returns two parts. `public` is aggregates only (counts, rates, topics): no review text, no names. It is what is saved with the
report and shown in the PDF and the monthly pack. `private` holds the unanswered critical reviews with their text and a draft reply: it lives in the visitor's session only
and is never saved. Reviewer names are never read at all.

Honesty: every figure is "of the N reviews in the file supplied", never a market figure or the hotel's overall reply rate on every platform. The reading of how replies were
written is rule-based and says so.
"""

import csv
import datetime as dt
import io
import re
import statistics

SOURCE = "Review pack (uploaded by the hotel)"
MAX_ROWS = 5000
CRITICAL_MAX = 2.5            # on a 5-point scale
POSITIVE_MIN = 4.0

# ----------------------------------------------------------------------- reading the file

_ALIASES = {
    "platform": ("platform", "source", "channel", "site", "ota", "reviewsite"),
    "date": ("reviewdate", "date", "created", "createdat", "createtime", "posted", "postdate", "submitted", "published", "publisheddate", "dateofreview", "reviewcreatedat", "reviewedon"),
    "rating": ("rating", "score", "stars", "starrating", "overallrating", "reviewscore", "guestrating", "overall", "ratingvalue", "reviewrating"),
    "text": ("review", "reviewtext", "comment", "comments", "text", "content", "body", "feedback", "reviewbody", "guestcomment", "reviewcomment"),
    "positive": ("positive", "positivereview", "pros", "liked", "whatyouliked", "positivecomment"),
    "negative": ("negative", "negativereview", "cons", "disliked", "whatyoudisliked", "negativecomment"),
    "title": ("title", "reviewtitle", "headline", "subject"),
    "reply": ("response", "reply", "ownerresponse", "managementresponse", "hotelreply", "hotelresponse", "reviewreplycomment", "replytext", "responsetext", "responsefromowner",
              "propertyreply", "propertyresponse", "managerresponse", "ownerreply", "replycomment"),
    "reply_date": ("responsedate", "replydate", "reviewreplyupdatetime", "dateofresponse", "repliedat", "respondedon", "responsetime", "replytime", "responsecreated"),
}
_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}


def _key(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def read_table(data, filename=""):
    """bytes of a CSV or Excel file -> (list of {header: value}, error). The first row that looks like a header is used."""
    name = (filename or "").lower()
    rows = []
    try:
        if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        else:
            text = None
            for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
                try:
                    text = data.decode(enc)
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue
            if text is None:
                return [], "The file could not be read as text."
            try:
                dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    except Exception as e:  # noqa: BLE001 - a bad file is a message, not a crash
        return [], f"The file could not be opened ({type(e).__name__}). Use a .csv or .xlsx export."
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if len(rows) < 2:
        return [], "The file has no reviews in it."
    head_i = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if _key(c) in {a for al in _ALIASES.values() for a in al}) >= 2), 0)
    headers = [str(c or "").strip() for c in rows[head_i]]
    out = []
    for r in rows[head_i + 1:head_i + 1 + MAX_ROWS]:
        out.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers) if h})
    return out, ""


def map_columns(headers):
    """{field: header} for the fields we can find. Needs at least a rating or a review text."""
    by_key = {_key(h): h for h in headers}
    m = {}
    for field, aliases in _ALIASES.items():
        for a in aliases:
            if a in by_key:
                m[field] = by_key[a]
                break
    return m


def parse_rating(v):
    """-> (value, scale) or (None, None). Handles 4, '4.5', '8,3', 'FOUR', '4 out of 5', '9/10'."""
    if v is None or str(v).strip() == "":
        return None, None
    s = str(v).strip().lower()
    if s in _WORDS:
        return float(_WORDS[s]), 5
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*(?:/|out of)\s*(\d+)", s)
    if m:
        return float(m.group(1).replace(",", ".")), int(m.group(2))
    m = re.search(r"\d+(?:[.,]\d+)?", s)
    if not m:
        return None, None
    return float(m.group(0).replace(",", ".")), None


def parse_date(v, dayfirst=True):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:          # an Excel serial date
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = str(v).strip()
    s = re.sub(r"(?<=\d)(st|nd|rd|th)\b", "", s)
    fmts = ["%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%d-%b-%Y", "%B %Y", "%b %Y"]
    fmts += ["%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M"] if dayfirst else ["%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%m/%d/%Y %H:%M"]
    s2 = re.sub(r"\.\d+", "", s)
    for f in fmts:
        try:
            return dt.datetime.strptime(s2, f).date()
        except ValueError:
            continue
    return None


def _dayfirst(rows, col):
    first_over, second_over = False, False
    for r in rows:
        m = re.match(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-]\d{2,4}", str(r.get(col) or ""))
        if m:
            first_over |= int(m.group(1)) > 12
            second_over |= int(m.group(2)) > 12
    return not (second_over and not first_over)


def normalise(rows, mapping, default_platform=""):
    """The reviews as plain dicts: {platform, date, rating (0-5 or None), text, positive, negative, reply, reply_date}. Names and any other column are never read."""
    scale_max = 0.0
    parsed = []
    for r in rows:
        v, sc = parse_rating(r.get(mapping.get("rating")))
        parsed.append((v, sc))
        if v is not None:
            scale_max = max(scale_max, sc or v)
    ten = scale_max > 5.01
    dcol, rcol = mapping.get("date"), mapping.get("reply_date")
    dfirst = _dayfirst(rows, dcol) if dcol else True
    rfirst = _dayfirst(rows, rcol) if rcol else True
    out = []
    for r, (v, sc) in zip(rows, parsed):
        rating = None
        if v is not None:
            rating = round(min(max(v * 5 / (sc or (10 if ten else 5)), 0), 5), 2)
        txt = " ".join(str(r.get(mapping[k]) or "").strip() for k in ("title", "text") if mapping.get(k) and k in mapping).strip()
        pos = str(r.get(mapping.get("positive")) or "").strip() if mapping.get("positive") else ""
        neg = str(r.get(mapping.get("negative")) or "").strip() if mapping.get("negative") else ""
        if not txt and (pos or neg):
            txt = (pos + " " + neg).strip()
        plat = str(r.get(mapping.get("platform")) or "").strip() if mapping.get("platform") else ""
        out.append({"platform": (plat or default_platform or "Not stated")[:30], "date": parse_date(r.get(dcol), dfirst) if dcol else None, "rating": rating, "text": txt, "positive": pos,
                    "negative": neg, "reply": str(r.get(mapping.get("reply")) or "").strip() if mapping.get("reply") else "",
                    "reply_date": parse_date(r.get(rcol), rfirst) if rcol else None})
    return [x for x in out if x["rating"] is not None or x["text"] or x["reply"]]


# ----------------------------------------------------------------------- reading replies

_STOP = set("""the and for that with this have from they were been very there their would about which when what your you our was are but not had has she him her his its
 also just them than then only into over some such will more most much many here well after before while where hotel stay stayed room rooms staff guest guests thank thanks
 dear regards kind best team manager general wishes please sorry again visit hope welcome feedback review reviews""".split())
_ACK = re.compile(r"\b(thank(?:s| you)?|apolog\w+|sorry|regret\w*|appreciate\w*|grateful)\b", re.I)
_ACTION = re.compile(r"\b(contact (?:us|me|our)|get in touch|reach out|email (?:us|me)|call (?:us|me)|please (?:contact|email|call|write)|speak (?:with|to) (?:you|us)|discuss|look(?:ed|ing)? into|investigat\w+|"
                     r"(?:have|has) (?:now |since |already )?(?:been )?(?:resolved|addressed|improved|replaced|renovated|refurbished|fixed|taken)|shared (?:this|it|your (?:feedback|comments))\b|"
                     r"taken (?:this|your) (?:on board|feedback)|made changes|will (?:be )?(?:improv\w+|address\w+|review\w*|look\w*|ensure|make sure)|put (?:this|it) right)\b", re.I)
_SIGN = re.compile(r"(?:regards|sincerely|best wishes|warm wishes|warmly|yours|many thanks|thank you|with thanks|cheers)[,.!\s]*(?:\n|\s)*[A-Z][\w'.-]*(?:\s+[A-Z][\w'.-]*){0,4}\s*$|"
                   r"\b(?:general manager|hotel manager|guest relations|front office|reception|duty manager|director|team|gm|rooms division|guest experience|management)\b\W*$", re.I)


def _words(t):
    return {w for w in re.findall(r"[a-z]{5,}", (t or "").lower()) if w not in _STOP}


def reply_flags(review_text, reply, hotel_words=()):
    """The four judgements about one reply, each True/False: personal (refers to specifics of the stay), acknowledges, offers_action, signed."""
    rw = _words(review_text) - set(hotel_words)
    shared = rw & _words(reply)
    personal = len(shared) >= 2 or (len(shared) >= 1 and len(rw) <= 4)             # a short review has few words to echo
    tail = (reply or "").strip()[-140:]
    return {"personal": personal, "acknowledges": bool(_ACK.search(reply or "")), "offers_action": bool(_ACTION.search(reply or "")), "signed": bool(_SIGN.search(tail))}


def _normal_reply(t):
    t = re.sub(r"^\s*(?:dear|hi|hello|hello and thank you|good (?:morning|afternoon|evening))[^\n,.!]{0,40}[,!.\n]", " ", (t or "").lower())
    t = re.sub(r"(?:kind |best |warm |many )?(?:regards|wishes|thanks|sincerely)[\s\S]{0,80}$", " ", t)
    return re.sub(r"[^a-z ]", " ", t)


def _shingles(t):
    w = t.split()
    return {" ".join(w[i:i + 3]) for i in range(max(len(w) - 2, 1))}


def template_groups(replies, threshold=0.8):
    """Group replies that are identical or near-identical once the greeting and sign-off are removed. -> list of group sizes, largest first (singletons included)."""
    reps = []
    for r in replies:
        n = _normal_reply(r)
        if len(n.split()) >= 5:
            reps.append(_shingles(n))
    groups = []                                     # each: [representative shingles, size]
    for sh in reps:
        for g in groups:
            u = len(sh | g[0])
            if u and len(sh & g[0]) / u >= threshold:
                g[1] += 1
                break
        else:
            groups.append([sh, 1])
    return sorted((g[1] for g in groups), reverse=True)


# ----------------------------------------------------------------------- topics

TOPICS = [
    ("breakfast", r"breakfast|buffet|brunch"), ("cleanliness", r"\bclean\w*|dirt\w*|dust\w*|stain\w*|smell\w*|mould|mold|spotless|filthy|hair\b"),
    ("staff and service", r"\bstaff\b|reception|front desk|concierge|helpful|rude|friendly|service|manager"), ("location", r"\blocation\b|located|close to|walking distance|station|metro|city cent|airport|transport"),
    ("noise", r"\bnois\w*|\bloud\b|\bquiet\b|thin walls|soundproof|street noise|traffic"), ("parking", r"parking|car park|garage|valet"), ("wi-fi", r"wi-?fi|internet|wireless"),
    ("value and price", r"\bvalue\b|\bprice\w*|expensive|overpriced|cheap|worth|\bcost\w*|\bcharg\w*"), ("bed and comfort", r"\bbeds?\b|mattress|pillow|comfortable|comfy"),
    ("bathroom", r"bathroom|shower|toilet|bath\b|water pressure|towel"), ("air conditioning and heating", r"air[- ]?con\w*|\bac\b|heating|too hot|too cold|temperature|radiator"),
    ("check-in and check-out", r"check-?in|check-?out|arrival|queue|waiting|late check"), ("pool, spa and gym", r"\bpool\b|\bspa\b|\bgym\b|sauna|fitness"),
    ("food and bar", r"restaurant|dinner|\bbar\b|\bmenu\b|\bfood\b|drinks|lunch"), ("room condition", r"dated|tired|renovat\w*|refurb\w*|worn|modern|old[- ]fashioned"), ("lift", r"\blifts?\b|elevator"),
]
_TOPIC_RX = [(t, re.compile(p, re.I)) for t, p in TOPICS]
# which topics the hotel's own site is judged on, and by what: guest-question ids from guest_questions.py
SITE_KEY = {"breakfast": "breakfast", "parking": "parking", "wi-fi": "wifi", "check-in and check-out": "checkin_out", "location": "transport"}


def topics_of(text):
    return [t for t, rx in _TOPIC_RX if rx.search(text or "")]


def _count_topics(reviews):
    praised, criticised = {}, {}
    n_pos = n_neg = 0
    has_split = any(r["positive"] or r["negative"] for r in reviews)
    for r in reviews:
        if has_split:
            p, g = topics_of(r["positive"]), topics_of(r["negative"])
            n_pos += bool(r["positive"])
            n_neg += bool(r["negative"])
        elif r["rating"] is not None and r["rating"] >= POSITIVE_MIN:
            p, g = topics_of(r["text"]), []
            n_pos += 1
        elif r["rating"] is not None and r["rating"] <= CRITICAL_MAX + 0.5:
            p, g = [], topics_of(r["text"])
            n_neg += 1
        else:
            continue
        for t in p:
            praised[t] = praised.get(t, 0) + 1
        for t in g:
            criticised[t] = criticised.get(t, 0) + 1
    return praised, criticised, n_pos, n_neg


# ----------------------------------------------------------------------- the analysis

def _rate(n, d):
    return None if not d else round(100 * n / d)


def analyse(reviews, hotel="", site_states=None, today=None):
    """
    reviews: normalise() output. site_states: {topic: 'answered'|'partial'|'not_found'|None} for the topics the hotel's own site is judged on.
    -> {"public": aggregates (safe to save), "private": {"unanswered": [...]}} or None when there is nothing to analyse.
    """
    if not reviews:
        return None
    today = today or dt.date.today()
    dated = [r["date"] for r in reviews if r["date"]]
    last = max(dated) if dated else None
    hotel_words = {w for w in re.findall(r"[a-z]{4,}", (hotel or "").lower())}
    for r in reviews:
        r["replied"] = bool(r["reply"])
        r["critical"] = r["rating"] is not None and r["rating"] <= CRITICAL_MAX
        r["band"] = None if r["rating"] is None else "critical" if r["rating"] <= CRITICAL_MAX else "positive" if r["rating"] >= POSITIVE_MIN else "mixed"
    has_reply_col = any(r["replied"] for r in reviews)
    n = len(reviews)
    replied = [r for r in reviews if r["replied"]]

    def block(rs):
        return {"reviews": len(rs), "replied": sum(1 for r in rs if r["replied"]), "rate": _rate(sum(1 for r in rs if r["replied"]), len(rs))}
    by_platform = {}
    for r in reviews:
        by_platform.setdefault(r["platform"], []).append(r)
    by_band = {b: block([r for r in reviews if r["band"] == b]) for b in ("critical", "mixed", "positive")}
    recent = [r for r in reviews if last and r["date"] and (last - r["date"]).days <= 90]
    # response time
    days = [(r["reply_date"] - r["date"]).days for r in replied if r["date"] and r["reply_date"] and r["reply_date"] >= r["date"]]
    rt = {"measured": len(days), "median_days": (round(statistics.median(days), 1) if days else None), "within_3_days_pct": _rate(sum(1 for d in days if d <= 3), len(days)),
          "within_7_days_pct": _rate(sum(1 for d in days if d <= 7), len(days))}
    # how replies were written
    flags = [reply_flags(r["text"], r["reply"], hotel_words) for r in replied]
    quality = {k: {"yes": sum(1 for f in flags if f[k]), "judged": len(flags)} for k in ("personal", "acknowledges", "offers_action", "signed")}
    groups = template_groups([r["reply"] for r in replied])
    in_groups = sum(g for g in groups if g > 1)
    tmpl = {"replies_compared": sum(groups), "in_repeated_templates": in_groups, "pct": _rate(in_groups, sum(groups)), "largest_group": groups[0] if groups and groups[0] > 1 else 0}
    # topics
    praised, criticised, n_pos, n_neg = _count_topics(reviews)
    top = lambda d, base: [{"topic": k, "reviews": v, "pct": _rate(v, base)} for k, v in sorted(d.items(), key=lambda kv: (-kv[1], kv[0]))[:8]]      # noqa: E731
    praised_rows, crit_rows = top(praised, n_pos), top(criticised, n_neg)
    for row in praised_rows + crit_rows:
        key = SITE_KEY.get(row["topic"])
        row["site"] = (site_states or {}).get(key) if key else None
    # months
    months = {}
    for r in reviews:
        if r["date"] and r["rating"] is not None:
            months.setdefault(r["date"].strftime("%Y-%m"), []).append(r["rating"])
    trend = [{"month": m, "reviews": len(v), "avg_rating": round(sum(v) / len(v), 2)} for m, v in sorted(months.items())][-12:]
    rated = [r["rating"] for r in reviews if r["rating"] is not None]
    crit = [r for r in reviews if r["critical"]]
    unanswered_all = [r for r in crit if not r["replied"]]
    public = {
        "source": SOURCE, "reviews": n, "rated": len(rated), "has_replies_column": has_reply_col,
        "first_date": min(dated).isoformat() if dated else None, "last_date": last.isoformat() if last else None,
        "latest_review_age_days": (today - last).days if last else None,
        "platforms": {p: block(rs) for p, rs in sorted(by_platform.items(), key=lambda kv: -len(kv[1]))},
        "reply_rate": _rate(len(replied), n) if has_reply_col else None, "replied": len(replied),
        "by_rating": by_band, "recent_90": {**block(recent), "window_end": last.isoformat() if last else None} if recent else None,
        "avg_rating": round(sum(rated) / len(rated), 2) if rated else None,
        "critical": len(crit), "critical_replied": len(crit) - len(unanswered_all), "critical_unanswered": len(unanswered_all) if has_reply_col else None,
        "response_time": rt, "reply_quality": quality, "templates": tmpl, "praised": praised_rows, "criticised": crit_rows, "positive_reviews": n_pos, "critical_reviews": n_neg, "trend": trend,
        "note": f"Every figure is for the {n} review(s) in the file supplied, not for every review the hotel has on every platform. How replies were written is a rule-based reading.",
    }
    # the part that is never saved: the unanswered critical reviews, newest first, with text and a draft reply
    un = sorted(unanswered_all, key=lambda r: r["date"] or dt.date.min, reverse=True)[:15] if has_reply_col else []
    private = {"unanswered": [{"platform": r["platform"], "date": r["date"].isoformat() if r["date"] else "", "rating": r["rating"], "topics": topics_of(r["negative"] or r["text"])[:3],
                               "text": (r["negative"] or r["text"])[:300], "draft": draft_reply(topics_of(r["negative"] or r["text"])[:2], hotel)} for r in un]}
    return {"public": public, "private": private}


def draft_reply(topics, hotel=""):
    """A skeleton for a person to finish, never a finished reply: it states only what we know (the topics the review raised) and leaves the facts in [brackets]."""
    about = (" about the " + " and the ".join(topics)) if topics else ""
    return (f"Dear guest,\n\nThank you for taking the time to share your feedback{about}, and I am sorry that your stay with us did not meet your expectations.\n\n"
            "[One or two sentences naming what the guest raised and what you found out or have changed. Only say what is true.]\n\n"
            "I would welcome the chance to put this right. Please contact me directly at [email or phone] and mention [reference].\n\n"
            f"Kind regards,\n[Name], [Role], {hotel or '[Hotel]'}")


def apply(report, public):
    """
    Put an analysed review pack into a finished audit report: keep the aggregates in report['review_pack'], replace the not-assessed Reviews & Reputation category with the one
    worked out from the file, add the findings to the plan, and recompute the headline. Idempotent: applying a newer pack replaces the earlier one. Only aggregates go in.
    """
    import advice
    import consultant
    import headline
    import scoring
    report["review_pack"] = public
    base = report.get("scorecard")
    if base:
        report["scorecard"] = scoring.apply_review_pack(base, public)
    if report.get("consultant") and "recommendations" in report["consultant"]:
        consultant.add_recommendations(report["consultant"], advice._from_review_pack(public), "review_pack")
    headline.attach(report)
    return report


def site_states_from(guest):
    """{topic: state} for the topics the hotel's own site is judged on, from the audit's guest-question results."""
    by_id = {q["id"]: q["state"] for q in (guest or {}).get("questions", [])}
    return {t: by_id.get(k) for t, k in SITE_KEY.items() if by_id.get(k)}
