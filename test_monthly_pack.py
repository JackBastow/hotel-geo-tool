"""
Tests for the monthly Excel pack (monthly_pack.py): what each sheet holds, the change columns, and the honesty of the Read me.
Run: python test_monthly_pack.py
"""
import copy
import io
import json
import os
import tempfile

from openpyxl import load_workbook

import backend
import monthly_pack as mp

fails = []


def check(name, cond, detail=""):
    print(("  PASS  " if cond else "  FAIL  ") + name + ("" if cond else f"   {detail}"))
    if not cond:
        fails.append(name)


HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "tests_data", "sample_report.json"), encoding="utf8") as fh:
    REP = json.load(fh)
DB = os.path.join(tempfile.gettempdir(), "hotel_pack_test.db")
if os.path.exists(DB):
    os.remove(DB)
st = backend.open_store("sqlite:///" + DB.replace("\\", "/"))

# September and October runs of the same hotel, plus a second hotel in October
sep = copy.deepcopy(REP)
sep["meta"]["run_at"] = "2026-09-10T09:00:00+00:00"
sep["headline"]["score"] = 70
sep["ai_insights"]["scoreboard"]["share"] = 20
sep["ai_sample"]["summary"].update(generic_named=1, generic_asked=4, niche_named=0, niche_asked=2)
oct1 = copy.deepcopy(REP)
oct1["meta"]["run_at"] = "2026-10-05T09:00:00+00:00"
oct2 = copy.deepcopy(REP)
oct2["meta"]["run_at"] = "2026-10-20T09:00:00+00:00"
oct2["headline"]["score"] = 91
oct2["ai_insights"]["scoreboard"]["share"] = 40
oct2["ai_sample"]["summary"].update(generic_named=3, generic_asked=4, niche_named=2, niche_asked=2)
other = copy.deepcopy(REP)
other["meta"].update(hotel="Chain Hotel", city="Luxembourg", base="https://www.hilton.com/en/hotels/x-doubletree-luxembourg/", website="https://www.hilton.com/en/hotels/x-doubletree-luxembourg/",
                     run_at="2026-10-12T09:00:00+00:00")
for r in (sep, oct1, oct2, other):
    st.save_run(r)

in_oct = st.runs_in_months("2026-10", "2026-10")
check("a month's runs are found by month, oldest first", [r["id"] for r in in_oct] == [2, 4, 3] and len(st.runs_in_months("2026-09", "2026-10")) == 4)
data = mp.build_workbook([(r, st.get_report(r["id"])) for r in in_oct], st.list_runs(limit=1000))
wb = load_workbook(io.BytesIO(data))
check("the workbook has every sheet the dashboard and the reader need",
      wb.sheetnames == ["Read me", "Hotels", "Metrics", "Scores and why", "AI rankings", "AI competitors", "AI why", "Top 3 actions", "Sources"], wb.sheetnames)


def table(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]


print("1. Hotels: one row per audit")
h, hotels = table("Hotels")
check("one row per run in the month (3), with the hotel, city and month", len(hotels) == 3 and {r["Hotel"] for r in hotels} == {"Brooklands Hotel", "Chain Hotel"} and {r["Month"] for r in hotels} == {"2026-10"})
b = [r for r in hotels if r["Hotel"] == "Brooklands Hotel"]
check("a hotel audited twice in the month has both rows, and only the LAST is marked 'latest in month'", sorted(r["Latest in month"] for r in b) == ["No", "Yes"]
      and next(r for r in b if r["Latest in month"] == "Yes")["Readiness (0-100)"] == 91)
last = next(r for r in b if r["Latest in month"] == "Yes")
check("every scorecard category has its own column (blank when not assessed)", all(f"{label} (score)" in h for _, label in mp.CATEGORY_ORDER) and last["Reviews & Reputation (score)"] in (None, ""))
check("AI visibility: named and asked for both town-wide and niche, as counts and percentages, with the share of names and the place in the table",
      (last["Named, town-wide"], last["Asked, town-wide"], last["Named % town-wide"], last["Named, niche"], last["Named % niche"]) == (3, 4, 75, 2, 100) and last["Share of hotel names %"] == 40,
      {k: last[k] for k in ("Named, town-wide", "Asked, town-wide", "Named % town-wide", "Named, niche", "Named % niche", "Share of hotel names %")})
check("reviews, speed and the plan are on the row", last["Reviews examined"] == 3 and last["Hotel replied"] == 2 and last["Mobile speed"] == 41 and last["Recommendations in plan"] == 20)
check("the top three actions are written out in plain titles", all(last[k] for k in ("Top action 1", "Top action 2", "Top action 3")))

print("2. change since the hotel's latest run in the previous month")
check("the October latest run is compared with September: readiness 70 -> 91 = +21, share of names 20 -> 40 = +20 points, town-wide named 25% -> 75% = +50, niche 0% -> 100% = +100",
      (last["Readiness change"], last["Share of names change (pts)"], last["Named % town-wide change"], last["Named % niche change"]) == (21.0, 20.0, 50.0, 100.0),
      (last["Readiness change"], last["Share of names change (pts)"], last["Named % town-wide change"], last["Named % niche change"]))
check("the earlier October run is ALSO compared with September (a re-run in the month is never compared with its own month)", str(next(r for r in b if r["Latest in month"] == "No")["Previous run"]) == "2026-09-10")
chain = next(r for r in hotels if r["Hotel"] == "Chain Hotel")
check("a hotel's first-ever run has no previous run and blank changes (not zero)", chain["Previous run"] in (None, "") and chain["Readiness change"] in (None, ""))

print("3. Metrics: the tidy table for dashboards")
mh, metrics = table("Metrics")
check("one row per hotel, run and metric, with a group, a metric name, a numeric value and a unit", mh == ("Month", "Run id", "Run date", "Hotel", "City", "Hotel key", "Group", "Metric", "Value", "Unit"))
mm = {(r["Run id"], r["Metric"]): r["Value"] for r in metrics}
check("the headline numbers are there for the latest run (readiness 91, named % niche 100, share 40)", mm[(3, "Readiness")] == 91 and mm[(3, "Named % niche markets")] == 100 and mm[(3, "Share of hotel names")] == 40)
check("the eight readiness measures and the scorecard categories are separate groups", {"Scorecard category", "Readiness measure", "AI visibility", "Reviews", "Web", "Overall"} <= {r["Group"] for r in metrics})
check("a metric that was not assessed is ABSENT, never zero", not any(r["Value"] in ("", None) for r in metrics) and (3, "Reviews & Reputation") not in mm)

print("4. why")
sh, scores = table("Scores and why")
tys = {r["Type"] for r in scores}
check("overall, categories and readiness measures each get a row with a reason", {"Overall readiness", "Scorecard category", "Readiness measure"} <= tys and all(r["Why"] for r in scores if r["Type"] == "Readiness measure"))
cat = next(r for r in scores if r["Run id"] == 3 and r["Name"] == "OTAs & Travel Platforms")
check("a scored category carries its status and how strong the evidence is", cat["Score"] == 100 and cat["Evidence strength"] in ("low", "medium", "high"))
na = next(r for r in scores if r["Run id"] == 3 and r["Name"] == "Reviews & Reputation")
check("a category that was not assessed says so, with the reason, and has no score", na["Status / band"] == "not assessed" and na["Score"] in (None, "") and na["Why"])
ah, why = table("AI why")
check("the hypotheses are labelled as to be checked, not causes, with their evidence ids", "not a cause" in ah[5] and any(r["Type"] == "Hypothesis" and r["Evidence ids"] for r in why))

print("5. rankings, competitors, actions, sources")
rh, ranks = table("AI rankings")
r3 = [r for r in ranks if r["Run id"] == 3]
check("question by question: type (town-wide / niche / about the hotel), whether it was named and where", {"Town-wide", "Niche market", "About the hotel"} <= {r["Question type"] for r in r3}
      and any(r["Question type"] == "Niche market" and r["Hotel named"] == "Yes" and r["Position in list"] == 2 for r in r3))
check("niche rows say why the question was asked", any(r["Question type"] == "Niche market" and r["Why we asked it"] for r in r3))
ch, comps = table("AI competitors")
check("who the assistant named instead, per hotel, with the hotel itself marked, for both lists", {r["List"] for r in comps} == {"Town-wide", "Niche market"}
      and any(str(r["Hotel named by the assistant"]).startswith("THIS HOTEL: ") for r in comps) and any(r["Hotel named by the assistant"] == "Oatlands Park Hotel" for r in comps))
th, acts = table("Top 3 actions")
per_run = {}
for r in acts:
    per_run.setdefault(r["Run id"], []).append(r["Rank"])
check("exactly the top three actions per run, ranked 1-3, with why and what to do", all(v == [1, 2, 3] for v in per_run.values()) and len(per_run) == 3 and all(r["Why it matters"] and r["What to do"] and r["Owner"] for r in acts))
sh2, srcs = table("Sources")
check("sources show whether the hotel was found there", any(r["Site the assistant used"] == "booking.com" and r["Hotel found there"] == "Yes" for r in srcs))

print("6. the Read me and the workbook itself")
text = " ".join(str(r[0]) for r in wb["Read me"].iter_rows(values_only=True) if r[0])
check("the Read me states the caveats: a sample of one model, readiness is an indicator, blanks are not zero", all(w in text for w in ("SAMPLE", "one AI model", "indicator", "Blank cells mean not assessed")))
ws = wb["Hotels"]
check("headers are frozen and filterable, so it opens ready to use", ws.freeze_panes == "A2" and ws.auto_filter.ref)
empty = load_workbook(io.BytesIO(mp.build_workbook([], [])))
check("a month with no runs still produces a valid, empty workbook", "Hotels" in empty.sheetnames and empty["Hotels"].max_row == 1)
check("the filename says which months it covers", mp.filename("2026-10", "2026-10") == "hotel-audit-pack-2026-10.xlsx" and mp.filename("2026-09", "2026-10") == "hotel-audit-pack-2026-09-to-2026-10.xlsx")

print()
print("FAILURES:", fails if fails else "none")
raise SystemExit(1 if fails else 0)
